"""
Презентация: Datapipe UI, ключевые точки коров (живой walkthrough).

Стиль как docs/datapipe и docs/e2e-korovas.

Сборка из корня data-collector:
    python docs/datapipe/generate_cow_keypoints_walkthrough.py

Выход:
    docs/datapipe/Datapipe-Cow-Keypoints-Walkthrough.pptx
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent.parent
BUSINESS = ROOT.parent / "business"
IMG = ROOT / "img" / "cow-keypoints-walkthrough"
UI = IMG / "ui"
GT = IMG / "gt"
TRAIN = IMG / "train"
CVAT = IMG / "cvat"
LOGO = REPO / "e8-team-logo-1024.png"
OUT = ROOT / "Datapipe-Cow-Keypoints-Walkthrough.pptx"

spec = importlib.util.spec_from_file_location(
    "gen_training", BUSINESS / "generate_training_presentation.py"
)
gen = importlib.util.module_from_spec(spec)
sys.modules["gen_training"] = gen
spec.loader.exec_module(gen)

C = gen.C
W = gen.W
TOTAL = 15

bg = gen.bg
title = gen.title
text = gen.text
rect = gen.rect
oval = gen.oval
pill = gen.pill
add_logo = gen.add_logo
new_deck = gen.new_deck

MUTED_DARK = RGBColor(0xAE, 0xB7, 0xC7)
BORDER_DARK = RGBColor(0x2D, 0x36, 0x48)

# Единая сетка
ML = Inches(0.65)          # левый край
MR = Inches(0.65)          # правый отступ
CW = W - ML - MR           # ширина контента ≈ 12.03"
CONTENT_TOP = Inches(2.05)
MEDIA_TOP = Inches(1.95)
MEDIA_H = Inches(4.75)
FOOTER_Y = Inches(7.08)


def footer(slide, n: int, *, dark: bool = False) -> None:
    color = RGBColor(0x7C, 0x83, 0x90) if dark else C.faint
    text(
        slide,
        ML,
        FOOTER_Y,
        Inches(8.0),
        Inches(0.22),
        "Epoch8 · Datapipe UI · ключевые точки",
        size=8,
        color=color,
    )
    text(
        slide,
        Inches(12.15),
        FOOTER_Y,
        Inches(0.7),
        Inches(0.22),
        f"{n:02d}/{TOTAL:02d}",
        size=8,
        color=color,
        align=PP_ALIGN.RIGHT,
    )


def _fit_size(iw: int, ih: int, max_w, max_h) -> tuple[int, int]:
    if iw <= 0 or ih <= 0:
        return int(max_w), int(max_h)
    scale = min(float(max_w) / iw, float(max_h) / ih)
    return int(iw * scale), int(ih * scale)


def add_media(
    slide,
    path: Path | None,
    l,
    t,
    max_w,
    max_h,
    *,
    caption: str | None = None,
    accent=C.mint,
    theme: str = "dark",
    frame: bool = True,
    placeholder: str | None = None,
    valign: str = "top",
):
    """Вписывает картинку в бокс. Подпись цепляется к низу самой картинки."""
    fill = C.night_2 if theme == "dark" else C.white
    border = C.line if theme == "light" else BORDER_DARK

    if path is None or not path.exists():
        rect(slide, l, t, max_w, max_h, fill, border, radius=True)
        text(
            slide,
            l + Inches(0.2),
            t + max_h / 2 - Inches(0.35),
            max_w - Inches(0.4),
            Inches(0.7),
            placeholder or "Скриншот: добавить позже",
            size=11,
            color=C.muted,
            align=PP_ALIGN.CENTER,
        )
        return None

    from PIL import Image

    with Image.open(path) as img:
        iw, ih = img.size

    final_w, final_h = _fit_size(iw, ih, max_w, max_h)
    final_l = l + int((max_w - final_w) / 2)
    if valign == "middle":
        final_t = t + int((max_h - final_h) / 2)
    else:
        final_t = t

    if frame:
        pad = Inches(0.05)
        rect(
            slide,
            final_l - pad,
            final_t - pad,
            final_w + pad * 2,
            final_h + pad * 2,
            fill,
            border,
            radius=True,
        )
    pic = slide.shapes.add_picture(str(path), final_l, final_t, width=final_w, height=final_h)
    if caption:
        gen.add_caption(
            slide,
            final_l,
            final_t + final_h - Inches(0.14),
            final_w,
            caption,
            accent,
            theme=theme,
        )
    return pic


def point_row(slide, y, num: str, head: str, accent, *, dark: bool = True, x=None, w=4.4):
    x = ML if x is None else x
    fg = C.white if dark else C.ink
    pill(slide, x, y, Inches(0.42), num, accent, C.ink if accent == C.lime else C.white)
    text(slide, x + Inches(0.58), y + Inches(0.01), Inches(w), Inches(0.38), head, size=14, color=fg)
    return y + Inches(0.58)


def first_existing(*paths: Path) -> Path | None:
    for p in paths:
        if p.exists():
            return p
    return None


def left_points(slide, items, *, dark: bool = True, top=None, gap=0.42, w=4.35):
    """Ровный столбик пунктов слева, с небольшим отступом сверху."""
    y = top if top is not None else CONTENT_TOP + Inches(0.25)
    for num, head, accent in items:
        y = point_row(slide, y, num, head, accent, dark=dark, w=w)
        y += Inches(gap)
    return y


def split_media(slide, path, *, caption: str, accent, theme: str = "dark", frame: bool = False):
    """Правая колонка под скрин (единый размер на всех split-слайдах)."""
    add_media(
        slide,
        path,
        Inches(5.35),
        MEDIA_TOP,
        Inches(7.35),
        MEDIA_H,
        caption=caption,
        accent=accent,
        theme=theme,
        frame=frame,
    )


def full_media(slide, path, *, caption: str, accent, theme: str = "dark"):
    """Широкий скрин на весь контент."""
    add_media(
        slide,
        path,
        ML,
        MEDIA_TOP,
        CW,
        MEDIA_H,
        caption=caption,
        accent=accent,
        theme=theme,
        frame=False,
    )


# ─── слайды ───────────────────────────────────────────────────────────────────


def slide_cover(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.e8_black)
    rect(s, ML, Inches(1.72), Inches(2.6), Pt(3), C.e8_coral)
    oval(s, Inches(11.95), Inches(0.5), Inches(0.75), Inches(0.75), C.e8_teal)
    rect(s, Inches(0), Inches(6.35), Inches(4.2), Inches(1.15), C.e8_teal)
    rect(s, Inches(4.2), Inches(6.35), W - Inches(4.2), Inches(1.15), C.e8_coral)

    add_logo(s, LOGO, ML, Inches(0.5), Inches(2.4), Inches(1.1))
    text(
        s,
        ML,
        Inches(2.05),
        Inches(4.9),
        Inches(1.5),
        "Ключевые точки\nв Datapipe UI",
        size=34,
        bold=True,
        color=C.white,
    )
    text(
        s,
        ML,
        Inches(3.7),
        Inches(4.7),
        Inches(1.0),
        "От разметки в CVAT до обучения модели\nв одном интерфейсе",
        size=14,
        color=C.e8_coral_muted,
    )
    text(
        s,
        ML,
        Inches(5.0),
        Inches(4.7),
        Inches(0.35),
        "разметка · подготовка · обучение · метрики",
        size=12,
        color=C.e8_teal_muted,
    )

    add_media(
        s,
        first_existing(CVAT / "job.png", GT / "hero-collage.jpg", UI / "01-overview.png"),
        Inches(5.55),
        Inches(1.2),
        Inches(7.1),
        Inches(4.55),
        caption="Разметка в CVAT",
        accent=C.e8_teal,
        theme="dark",
        frame=False,
    )
    text(
        s,
        Inches(0.75),
        Inches(6.58),
        Inches(12.0),
        Inches(0.4),
        "Учебный стенд Korovas · Epoch8",
        size=17,
        bold=True,
        color=C.e8_black,
        align=PP_ALIGN.CENTER,
    )


def slide_agenda(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "Маршрут", "Что происходит с данными на каждом этапе", dark=True)

    items = [
        ("01", "CVAT", "Где лежит размеченный материал", C.coral),
        ("02", "Аннотация", "Забираем разметку и собираем датасет", C.mint),
        ("03", "Datapipe UI", "Смотрим данные, запуски и статусы", C.cobalt),
        ("04", "Обучение", "Заморозка, train, метрики", C.lime),
    ]
    gap = Inches(0.22)
    card_w = (CW - gap * 3) / 4
    x = ML
    for num, head, body, accent in items:
        rect(s, x, CONTENT_TOP, card_w, Inches(4.4), C.night_2, BORDER_DARK, radius=True)
        rect(s, x, CONTENT_TOP, Inches(0.1), Inches(4.4), accent)
        text(s, x + Inches(0.28), Inches(2.4), card_w - Inches(0.5), Inches(0.45), num, size=24, bold=True, color=accent)
        text(s, x + Inches(0.28), Inches(3.1), card_w - Inches(0.5), Inches(0.7), head, size=17, bold=True, color=C.white)
        text(s, x + Inches(0.28), Inches(3.95), card_w - Inches(0.5), Inches(1.9), body, size=13, color=MUTED_DARK)
        x += card_w + gap

    footer(s, n, dark=True)


def slide_cvat(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "CVAT: источник разметки", "Проекты с кадрами коров, боксами и ключевыми точками", dark=True)

    left_points(
        s,
        [
            ("1", "Проекты с пакетами кадров", C.coral),
            ("2", "Задачи и задания на разметку", C.mint),
            ("3", "На кадре: корова, бокс, точки", C.cobalt),
        ],
        dark=True,
        top=Inches(2.45),
        gap=0.5,
    )
    split_media(
        s,
        first_existing(CVAT / "project.png", CVAT / "task.png"),
        caption="Проекты в CVAT",
        accent=C.coral,
        theme="dark",
    )
    footer(s, n, dark=True)


def slide_cvat_job(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "Как выглядит разметка", "Бокс коровы и ключевые точки по схеме разметки")

    add_media(
        s,
        CVAT / "job.png",
        ML,
        MEDIA_TOP,
        Inches(8.15),
        MEDIA_H,
        caption="Рабочее место разметчика",
        accent=C.coral,
        theme="light",
        frame=False,
    )

    panel_l = Inches(9.05)
    panel_w = Inches(3.65)
    rect(s, panel_l, MEDIA_TOP, panel_w, MEDIA_H, C.white, C.line, radius=True)
    rect(s, panel_l, MEDIA_TOP, Inches(0.1), MEDIA_H, C.coral)
    text(s, panel_l + Inches(0.3), Inches(2.2), panel_w - Inches(0.5), Inches(0.4), "На кадре", size=16, bold=True, color=C.ink)

    bullets = [
        ("1", "Бокс коровы", C.coral),
        ("2", "Ключевые точки", C.mint),
        ("3", "Готово к пайплайну", C.cobalt),
    ]
    y = Inches(2.9)
    for num, head, accent in bullets:
        pill(s, panel_l + Inches(0.3), y, Inches(0.38), num, accent, C.white)
        text(s, panel_l + Inches(0.85), y + Inches(0.01), panel_w - Inches(1.15), Inches(0.38), head, size=13, color=C.ink)
        y += Inches(0.7)

    text(
        s,
        panel_l + Inches(0.3),
        Inches(5.3),
        panel_w - Inches(0.55),
        Inches(1.0),
        "Холка, маклок,\nседловидный бугор\nи другие точки схемы",
        size=12,
        color=C.muted,
    )
    footer(s, n)


def slide_into_pipeline(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "Разметка попадает в пайплайн", "Стадия аннотации забирает готовые задания из CVAT", dark=True)
    full_media(
        s,
        UI / "02-graph-annotation.png",
        caption="Запуск аннотации в Datapipe UI",
        accent=C.mint,
    )
    footer(s, n, dark=True)


def slide_images(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "Откуда берутся картинки", "Кадры могут лежать в разных местах")

    left_points(
        s,
        [
            ("1", "В облачном хранилище (S3 / MinIO)", C.cobalt),
            ("2", "Прямо в CVAT, как кадры задачи", C.mint),
            ("3", "Локально на диске, если уже скачали", C.coral),
        ],
        dark=False,
        top=Inches(2.4),
        gap=0.48,
        w=4.5,
    )
    text(
        s,
        ML,
        Inches(5.55),
        Inches(4.5),
        Inches(1.0),
        "Пайплайн сам понимает источник и подтягивает кадры туда, где они нужны.",
        size=13,
        color=C.muted,
    )
    split_media(
        s,
        UI / "05-images.png",
        caption="Картинки в Datapipe UI",
        accent=C.cobalt,
        theme="light",
        frame=True,
    )
    footer(s, n)


def slide_ground_truth(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "Ground truth", "Уже размеченные кадры: боксы и ключевые точки", dark=True)

    text(
        s,
        ML,
        Inches(1.95),
        CW,
        Inches(0.4),
        "В рабочем контуре разметка уже есть. Пайплайн берёт готовые аннотации как эталон.",
        size=14,
        color=MUTED_DARK,
    )

    gap = Inches(0.22)
    card_w = (CW - gap * 2) / 3
    paths = [GT / "01.jpg", GT / "02.jpg", GT / "04.jpg"]
    accents = [C.coral, C.mint, C.cobalt]
    x = ML
    for path, accent in zip(paths, accents):
        add_media(
            s,
            path,
            x,
            Inches(2.55),
            card_w,
            Inches(3.95),
            caption="Эталон",
            theme="dark",
            frame=False,
            accent=accent,
            valign="middle",
        )
        x += card_w + gap

    footer(s, n, dark=True)


def slide_split(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "Сплит и теги", "Как делим данные и помечаем группы", dark=True)

    cards = [
        ("Сплит", "Делим кадры на обучение, проверку и тест.\n\nМодель учится на одном наборе, а качество смотрим на другом.", C.cobalt),
        ("Теги", "Помечаем группы кадров, например по проекту или сценарию.\n\nПотом можно смотреть метрики отдельно по каждой группе.", C.mint),
    ]
    gap = Inches(0.28)
    card_w = (CW - gap) / 2
    x = ML
    for head, body, accent in cards:
        rect(s, x, CONTENT_TOP, card_w, Inches(4.35), C.night_2, BORDER_DARK, radius=True)
        rect(s, x, CONTENT_TOP, Inches(0.1), Inches(4.35), accent)
        text(s, x + Inches(0.45), Inches(2.45), card_w - Inches(0.8), Inches(0.55), head, size=22, bold=True, color=C.white)
        text(s, x + Inches(0.45), Inches(3.25), card_w - Inches(0.8), Inches(2.6), body, size=15, color=MUTED_DARK)
        x += card_w + gap

    footer(s, n, dark=True)


def slide_ui_overview(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "Datapipe UI", "Весь путь: данные, запуски, обучение", dark=True)

    left_points(
        s,
        [
            ("1", "Обзор пайплайна", C.cobalt),
            ("2", "История запусков", C.mint),
            ("3", "Статусы и логи на месте", C.lime),
        ],
        dark=True,
        top=Inches(2.5),
        gap=0.55,
        w=4.0,
    )
    split_media(
        s,
        UI / "01-overview.png",
        caption="Обзор в Datapipe UI",
        accent=C.mint,
        theme="dark",
    )
    footer(s, n, dark=True)


def slide_freeze(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "Заморозка датасета", "Фиксируем снимок данных перед обучением")

    left_points(
        s,
        [
            ("1", "Сохраняем состав кадров и разметки", C.cobalt),
            ("2", "Модель учится на понятной версии данных", C.mint),
            ("3", "Можно вернуться к тому же снимку позже", C.coral),
        ],
        dark=False,
        top=Inches(2.5),
        gap=0.55,
        w=4.4,
    )
    split_media(
        s,
        UI / "06-frozen.png",
        caption="Замороженные наборы",
        accent=C.mint,
        theme="light",
        frame=True,
    )
    footer(s, n)


def slide_train_start(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "Запуск обучения", "Из Datapipe UI: запрос на train и список запусков", dark=True)

    left_points(
        s,
        [
            ("1", "Выбираем замороженный набор", C.cobalt),
            ("2", "Запускаем обучение", C.mint),
            ("3", "Следим за статусом в запусках", C.coral),
        ],
        dark=True,
        top=Inches(2.5),
        gap=0.55,
        w=4.3,
    )
    split_media(
        s,
        UI / "03-runs.png",
        caption="Запуски",
        accent=C.coral,
        theme="dark",
    )
    footer(s, n, dark=True)


def slide_train_logs(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "Ход обучения", "Эпохи и метрики видны прямо в карточке запуска", dark=True)
    full_media(
        s,
        UI / "04-run-train-logs.png",
        caption="Запуск обучения в Datapipe UI",
        accent=C.lime,
    )
    footer(s, n, dark=True)


def slide_train_batches(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "Что видит модель", "Примеры батчей и контроль разметки")

    gap = Inches(0.22)
    card_w = (CW - gap * 2) / 3
    items = [
        (TRAIN / "batch0.jpg", "Батч", C.cobalt),
        (TRAIN / "batch1.jpg", "Батч", C.mint),
        (TRAIN / "labels.jpg", "Разметка", C.coral),
    ]
    x = ML
    for path, cap, accent in items:
        add_media(
            s,
            path,
            x,
            MEDIA_TOP,
            card_w,
            MEDIA_H,
            caption=cap,
            accent=accent,
            theme="light",
            frame=True,
            valign="middle",
        )
        x += card_w + gap

    footer(s, n)


def slide_metrics(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.night)
    title(s, "Метрики после обучения", "Качество модели в одной таблице", dark=True)
    full_media(
        s,
        UI / "08-metrics.png",
        caption="Метрики в Datapipe UI",
        accent=C.mint,
    )
    footer(s, n, dark=True)


def slide_close(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, C.e8_black)
    rect(s, ML, Inches(2.05), Inches(2.4), Pt(3), C.e8_coral)
    text(
        s,
        ML,
        Inches(2.35),
        CW,
        Inches(1.0),
        "Разметка → данные → модель",
        size=34,
        bold=True,
        color=C.white,
    )
    text(
        s,
        ML,
        Inches(3.55),
        Inches(11.5),
        Inches(0.9),
        "CVAT хранит разметку. Datapipe UI ведёт подготовку и обучение\nс понятной историей запусков и метрик.",
        size=15,
        color=C.e8_coral_muted,
    )
    flow = [
        ("CVAT", C.coral),
        ("аннотация", C.mint),
        ("заморозка", C.cobalt),
        ("обучение", C.lime),
        ("метрики", C.mint),
    ]
    gap = Inches(0.2)
    pill_w = (CW - gap * 4) / 5
    x = ML
    for label, color in flow:
        pill(s, x, Inches(5.15), pill_w, label, color, C.ink if color == C.lime else C.white)
        x += pill_w + gap
    footer(s, n, dark=True)


def main() -> None:
    prs = new_deck()
    builders = [
        slide_cover,
        slide_agenda,
        slide_cvat,
        slide_cvat_job,
        slide_into_pipeline,
        slide_images,
        slide_ground_truth,
        slide_split,
        slide_ui_overview,
        slide_freeze,
        slide_train_start,
        slide_train_logs,
        slide_train_batches,
        slide_metrics,
        slide_close,
    ]
    global TOTAL
    TOTAL = len(builders)
    for i, fn in enumerate(builders, 1):
        fn(prs, i)
    prs.save(OUT)
    print(f"Wrote {OUT} ({TOTAL} slides)")


if __name__ == "__main__":
    main()
